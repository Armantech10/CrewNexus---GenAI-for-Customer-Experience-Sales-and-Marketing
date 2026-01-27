from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# Create presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

def add_slide(prs, title, content_items, is_title_slide=False):
    slide_layout = prs.slide_layouts[6]  # Blank
    slide = prs.slides.add_slide(slide_layout)
    
    # Background (dark theme)
    background = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    background.fill.solid()
    background.fill.fore_color.rgb = RGBColor(15, 15, 25)
    background.line.fill.background()
    
    # Title
    title_top = Inches(2.5) if is_title_slide else Inches(0.4)
    title_box = slide.shapes.add_textbox(Inches(0.5), title_top, Inches(12), Inches(1))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(44) if is_title_slide else Pt(32)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER if is_title_slide else PP_ALIGN.LEFT
    
    # Content
    y_pos = Inches(3.8) if is_title_slide else Inches(1.6)
    for item in content_items:
        content_box = slide.shapes.add_textbox(Inches(0.6), y_pos, Inches(12), Inches(0.6))
        tf = content_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = item
        p.font.size = Pt(22) if is_title_slide else Pt(18)
        p.font.color.rgb = RGBColor(200, 200, 210)
        p.alignment = PP_ALIGN.CENTER if is_title_slide else PP_ALIGN.LEFT
        y_pos += Inches(0.55)
    
    return slide

# Slide 1: Title
add_slide(prs, "CrewNexus", [
    "GenAI for Customer Experience, Sales, and Marketing",
    "",
    "Team: [Your Team Name]",
    "Diptish De | [Teammate 2] | [Teammate 3] | [Teammate 4]",
    "[Your College/Organization]"
], is_title_slide=True)

# Slide 2: Overview
add_slide(prs, "The Problem: Fragmented Customer Experience", [
    "• Sales, Marketing, and Support teams operate in silos",
    "• Context is lost during customer handoffs between departments",
    "• Human teams cannot provide instant 24/7 personalized responses",
    "",
    "Our Vision:",
    "• A unified platform where specialized AI agents collaborate",
    "• Shared 'brain' delivers seamless, hyper-personalized CX"
])

# Slide 3: Data & Insights
add_slide(prs, "Supporting Data & Insights", [
    "• 76% of customers expect consistent interactions across departments",
    "   Yet 54% feel like they're talking to separate companies (Salesforce)",
    "",
    "• 35-50% of sales go to the vendor that responds first",
    "   Human avg response: 10+ hrs | AI response: <5 seconds",
    "",
    "• 20-30% revenue loss due to poor lead management & slow support",
    "",
    "Key Insight: The problem isn't automation—it's orchestration"
])

# Slide 4: Solution
add_slide(prs, "CrewNexus: The Solution", [
    "Architecture:",
    "• Unified Orchestrator: Analyzes intent & sentiment, routes to specialist",
    "• Multi-Agent Crew: Marketing (CrewAI), Sales (LangChain), Support (RAG)",
    "• Shared Memory: Redis short-term + Qdrant long-term vector memory",
    "",
    "Feasibility:",
    "• MVP Built: FastAPI backend + Next.js frontend (functional prototype)",
    "• Scalable: Docker containerized | Cost-Effective: Efficient LLM routing",
    "",
    "Impact: 360° Customer View | 24/7 Autonomy | Real-time Personalization"
])

# Save
output_path = r"d:\gen Ai\unified-genai-platform\docs\CrewNexus_Presentation.pptx"
prs.save(output_path)
print(f"Presentation saved to: {output_path}")
